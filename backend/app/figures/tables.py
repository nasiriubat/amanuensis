"""Results tables (ROADMAP day 11): experiment data the author uploads as CSV or XLSX.

results/index.json      [{name, title, caption, columns, rows, source_file, created_at, summary}]
results/<name>.csv      the data, normalised to UTF-8 comma-separated

A table is the only place a number in the Evaluation may come from besides the interview. Facts
extraction reads a summary of every table, and evaluation-type sections receive the tables as
Markdown so the draft can quote them. Nothing is inferred beyond what the cells say.
"""

from __future__ import annotations

import csv
import io
import json
import re
import statistics
from datetime import UTC, datetime
from pathlib import Path

from .. import storage

MAX_ROWS = 5000
MAX_COLS = 40
PREVIEW_ROWS = 30
PROMPT_CHARS = 4000
NAME_RE = re.compile(r"^[a-z0-9][a-z0-9\-]{0,40}$")


def _dir(root: Path) -> Path:
    d = root / "results"
    d.mkdir(exist_ok=True)
    return d


def _now() -> str:
    return datetime.now(UTC).isoformat()


def load(root: Path) -> list[dict]:
    p = _dir(root) / "index.json"
    if p.exists():
        try:
            return json.loads(p.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            pass
    return []


def save(root: Path, items: list[dict]) -> None:
    storage.write_text(_dir(root) / "index.json", json.dumps(items, indent=2, ensure_ascii=False))


def get(root: Path, name: str) -> dict | None:
    return next((t for t in load(root) if t["name"] == name), None)


def slug_name(name: str, taken: set[str]) -> str:
    base = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")[:40] or "table"
    out, n = base, 2
    while out in taken:
        out = f"{base[:36]}-{n}"
        n += 1
    return out


# ------------------------------------------------------------------ parsing


def parse_csv(data: bytes) -> tuple[list[str], list[list[str]]]:
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [r for r in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in r)]
    if not rows:
        raise ValueError("The file has no rows")
    header = [c.strip() or f"column {i + 1}" for i, c in enumerate(rows[0])]
    body = [[c.strip() for c in r] + [""] * (len(header) - len(r)) for r in rows[1:]]
    return header[:MAX_COLS], [r[:MAX_COLS] for r in body[:MAX_ROWS]]


def parse_xlsx(data: bytes) -> tuple[list[str], list[list[str]]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows: list[list[str]] = []
    for r in ws.iter_rows(values_only=True):
        cells = ["" if v is None else str(v).strip() for v in r]
        if any(cells):
            rows.append(cells)
        if len(rows) > MAX_ROWS:
            break
    if not rows:
        raise ValueError("The first sheet has no rows")
    header = [c or f"column {i + 1}" for i, c in enumerate(rows[0])]
    body = [r + [""] * (len(header) - len(r)) for r in rows[1:]]
    return header[:MAX_COLS], [r[:MAX_COLS] for r in body]


def parse(data: bytes, filename: str) -> tuple[list[str], list[list[str]]]:
    low = filename.lower()
    if low.endswith((".xlsx", ".xlsm")):
        return parse_xlsx(data)
    if low.endswith((".csv", ".tsv", ".txt")):
        return parse_csv(data)
    raise ValueError("Upload a .csv, .tsv or .xlsx file")


# ------------------------------------------------------------------ summaries


def _num(v: str) -> float | None:
    t = v.replace(",", "").replace("%", "").strip()
    if not t:
        return None
    try:
        return float(t)
    except ValueError:
        return None


def summarize(columns: list[str], rows: list[list[str]]) -> list[dict]:
    """Per column: numeric range and mean, or distinct values. Only what the cells say."""
    out = []
    for i, col in enumerate(columns):
        values = [r[i] for r in rows if i < len(r) and r[i] != ""]
        nums = [n for n in (_num(v) for v in values) if n is not None]
        if values and len(nums) >= max(1, int(0.8 * len(values))):
            out.append(
                {
                    "column": col,
                    "kind": "number",
                    "count": len(nums),
                    "min": round(min(nums), 3),
                    "max": round(max(nums), 3),
                    "mean": round(statistics.fmean(nums), 3),
                }
            )
        else:
            distinct = list(dict.fromkeys(values))
            out.append(
                {"column": col, "kind": "text", "count": len(values), "distinct": len(distinct), "top": distinct[:6]}
            )
    return out


def markdown_table(columns: list[str], rows: list[list[str]], max_rows: int = PREVIEW_ROWS) -> str:
    def esc(c: str) -> str:
        return c.replace("|", "\\|")

    lines = ["| " + " | ".join(esc(c) for c in columns) + " |", "|" + "---|" * len(columns)]
    for r in rows[:max_rows]:
        lines.append("| " + " | ".join(esc(c) for c in r[: len(columns)]) + " |")
    if len(rows) > max_rows:
        lines.append(f"| … {len(rows) - max_rows} more rows … |" + " |" * (len(columns) - 1))
    return "\n".join(lines)


def prompt_block(root: Path, max_chars: int = PROMPT_CHARS) -> str:
    """Every table as a compact text block for facts extraction and evaluation drafting."""
    items = load(root)
    if not items:
        return ""
    parts = []
    for t in items:
        cols, rows = read_rows(root, t["name"])
        lines = [f"### Table {t['name']}: {t.get('caption') or t.get('title') or ''} ({len(rows)} rows)"]
        for s in t.get("summary") or summarize(cols, rows):
            if s["kind"] == "number":
                lines.append(
                    f"- {s['column']}: numbers, n={s['count']}, min {s['min']}, max {s['max']}, mean {s['mean']}"
                )
            else:
                lines.append(f"- {s['column']}: {s['distinct']} distinct values, e.g. {', '.join(s['top'][:4])}")
        lines.append(markdown_table(cols, rows, max_rows=12))
        parts.append("\n".join(lines))
    text = "\n\n".join(parts)
    return text[:max_chars] + ("\n[… tables truncated to the budget …]" if len(text) > max_chars else "")


# ------------------------------------------------------------------ CRUD


def create(root: Path, filename: str, data: bytes, caption: str = "") -> dict:
    columns, rows = parse(data, filename)
    items = load(root)
    name = slug_name(Path(filename).stem, {t["name"] for t in items})
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(columns)
    w.writerows(rows)
    storage.write_text(_dir(root) / f"{name}.csv", buf.getvalue())
    item = {
        "name": name,
        "title": Path(filename).stem,
        "caption": caption.strip(),
        "columns": columns,
        "rows": len(rows),
        "source_file": filename,
        "created_at": _now(),
        "summary": summarize(columns, rows),
    }
    items.append(item)
    save(root, items)
    storage.git_commit(root, f"Results table: add {name}")
    return item


def read_rows(root: Path, name: str) -> tuple[list[str], list[list[str]]]:
    if not NAME_RE.match(name):
        raise ValueError("Bad table name")
    p = _dir(root) / f"{name}.csv"
    if not p.exists():
        raise ValueError("Table not found")
    rows = list(csv.reader(io.StringIO(p.read_text(encoding="utf-8"))))
    return (rows[0], rows[1:]) if rows else ([], [])


def update(root: Path, name: str, caption: str) -> dict:
    items = load(root)
    for t in items:
        if t["name"] == name:
            t["caption"] = caption.strip()
            t["updated_at"] = _now()
            save(root, items)
            return t
    raise ValueError("Table not found")


def delete(root: Path, name: str) -> bool:
    items = load(root)
    keep = [t for t in items if t["name"] != name]
    if len(keep) == len(items):
        return False
    save(root, keep)
    p = _dir(root) / f"{name}.csv"
    if NAME_RE.match(name) and p.exists():
        p.unlink()
    storage.git_commit(root, f"Results table: remove {name}")
    return True


def insert_snippet(root: Path, name: str) -> str:
    """What the Studio inserts: the table as Markdown with a pandoc caption line."""
    t = get(root, name)
    if not t:
        raise ValueError("Table not found")
    cols, rows = read_rows(root, name)
    caption = t.get("caption") or t.get("title") or name
    return f"\n\n{markdown_table(cols, rows, max_rows=PREVIEW_ROWS)}\n\nTable: {caption} {{#tbl:{name}}}\n\n"
